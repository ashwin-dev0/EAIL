"""Bounded extraction. Invoked in a separate process by src.ingest."""
import csv, io, json, sys
from pathlib import Path
from src.config import settings
from src.file_safety import inspect_bytes

def extract(path):
    ext=path.suffix.lower(); data=path.read_bytes()
    inspect_bytes(data,ext)
    sections=[]; total=0
    def add(label,text):
        nonlocal total
        text=str(text).strip()
        if not text: return
        total+=len(text)
        if total>settings.max_text or len(sections)>=settings.max_sections: raise ValueError('Extraction limit exceeded')
        sections.append((str(label),text))
    if ext in {'.txt','.md'}: add('Section 1',data.decode('utf-8-sig'))
    elif ext=='.json': add('JSON',json.dumps(json.loads(data),ensure_ascii=False,indent=2))
    elif ext=='.csv':
        reader=csv.reader(io.StringIO(data.decode('utf-8-sig')))
        header=next(reader,[])
        for i,row in enumerate(reader,1):
            if i>settings.max_rows: raise ValueError('CSV row limit')
            add(f'Row {i}', ' | '.join(f'{header[j] if j<len(header) else j}: {v}' for j,v in enumerate(row)))
    elif ext in {'.xml','.html','.htm'}:
        if ext=='.xml':
            from defusedxml import ElementTree as ET
            add('XML',' '.join(t.strip() for t in ET.fromstring(data).itertext() if t.strip()))
        else:
            from bs4 import BeautifulSoup
            soup=BeautifulSoup(data,'html.parser')
            for el in soup(['script','style','noscript']): el.decompose()
            add('HTML',soup.get_text(' ',strip=True))
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted: raise ValueError('Encrypted PDF rejected')
        if len(reader.pages)>settings.max_sections: raise ValueError('PDF page limit')
        for i,page in enumerate(reader.pages,1):
            text=page.extract_text() or ''
            if not text.strip():
                import pypdfium2 as pdfium
                doc=pdfium.PdfDocument(data)
                try:
                    image=doc[i-1].render(scale=1.5).to_pil()
                    text=ocr(image)
                finally: doc.close()
            add(f'Page {i}',text)
    elif ext in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'}:
        from PIL import Image, ImageSequence
        Image.MAX_IMAGE_PIXELS=25_000_000
        with Image.open(io.BytesIO(data)) as image:
            for i,frame in enumerate(ImageSequence.Iterator(image),1):
                if i>settings.max_sections: raise ValueError('Image frame limit')
                if frame.width*frame.height>25_000_000: raise ValueError('Image pixel limit')
                add(f'Image {i}',ocr(frame.copy()))
    elif ext=='.docx':
        from docx import Document
        doc=Document(io.BytesIO(data))
        text='\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows: text+='\n'+' | '.join(c.text for c in row.cells)
        add('Document',text)
    elif ext=='.pptx':
        from pptx import Presentation
        deck=Presentation(io.BytesIO(data))
        if len(deck.slides)>settings.max_sections: raise ValueError('Slide limit')
        for i,slide in enumerate(deck.slides,1):
            parts=[]
            for shape in slide.shapes:
                if shape.has_text_frame: parts.append(shape.text)
                if shape.has_table:
                    parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            add(f'Slide {i}','\n'.join(parts))
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        count=0
        try:
            for sheet in book:
                # Read actual rows, not worksheet-reported dimensions.
                for i,row in enumerate(sheet.iter_rows(values_only=True),1):
                    count+=1
                    if count>settings.max_rows: raise ValueError('Spreadsheet row limit')
                    if len(row)>200: raise ValueError('Spreadsheet column limit')
                    values=[str(v) if v is not None else '' for v in row]
                    add(f'{sheet.title} row {i}', ' | '.join(values))
        finally: book.close()
    elif ext in {'.odt','.ods','.odp'}:
        from odf.opendocument import load
        from odf import teletype
        from odf.text import P
        from odf.table import TableRow
        from odf.draw import Page
        doc=load(str(path))
        if ext=='.odt': add('Document','\n'.join(teletype.extractText(p) for p in doc.getElementsByType(P)))
        elif ext=='.ods':
            rows=doc.getElementsByType(TableRow)
            if len(rows)>settings.max_rows: raise ValueError('ODS row limit')
            for i,row in enumerate(rows,1): add(f'Row {i}',teletype.extractText(row))
        else:
            for i,page in enumerate(doc.getElementsByType(Page),1): add(f'Slide {i}',teletype.extractText(page))
    else: raise ValueError('Unsupported format')
    return sections

def ocr(image):
    import pytesseract
    return pytesseract.image_to_string(image,timeout=settings.ocr_timeout)

def main():
    # Parent imposes wall-clock timeout; process imposes CPU/address-space limits.
    if sys.platform.startswith('linux'):
        import resource
        resource.setrlimit(resource.RLIMIT_CPU,(settings.extraction_timeout,settings.extraction_timeout+1))
        cap=settings.extraction_memory_mb*1024*1024
        resource.setrlimit(resource.RLIMIT_AS,(cap,cap))
    try:
        sections=extract(Path(sys.argv[1]))
        Path(sys.argv[2]).write_text(json.dumps(sections,ensure_ascii=False))
    except Exception as exc:
        # No source contents or parser exception text leaked into logs/output.
        print(type(exc).__name__,file=sys.stderr);raise SystemExit(1)
if __name__=='__main__': main()
